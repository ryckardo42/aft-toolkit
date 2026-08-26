# -*- coding: utf-8 -*-
"""
Monta o Relatorio de Analise de Acidente (.docx) a partir de um arquivo JSON de conteudo.

Uso:
    python gerar_relatorio_docx.py conteudo.json

O JSON descreve o relatorio; este script cuida da formatacao (margens, fonte, cores,
tabela de identificacao, secoes, subtitulos, bullets e blocos de fator SFIT). O
cabecalho institucional - brasao, Ministerio do Trabalho e Emprego, Secretaria de
Inspecao do Trabalho, lotacao do AFT e logos SIT/AFT - vem do _scripts/cabecalho.py,
o mesmo dos demais .docx do toolkit.
A skill /aft-analise-acidente preenche o conteudo; a formatacao fica padronizada aqui.

Esquema do JSON:
{
  "saida": "caminho\\\\Relatorio.docx",          # obrigatorio
  "titulo": "RELATORIO DE ANALISE DE ACIDENTE DO TRABALHO",
  "subtitulo": "Acidente do trabalho tipico com obito",   # opcional
  "identificacao": [["Empregador","..."], ["CNPJ","..."]],# linhas da tabela de capa
  "sumario": true,                                        # opcional: indice com paginas
  "secoes": [
     {"titulo": "1. DESCRICAO DO LOCAL DO ACIDENTE",
      "blocos": [
         {"t":"p",   "x":"paragrafo, aceita **negrito** inline"},
         {"t":"sub", "x":"1.1 Subtopico de 2o nivel"},
         {"t":"sub3","x":"5.4.1 Subtopico de 3o nivel"},
         {"t":"b",   "x":"item de lista (bullet)"},
         {"t":"fator","codigo":"251008","nome":"...","classe":"determinante","desc":"..."}
      ]}
  ],
  "rodape": "texto final em italico pequeno"             # opcional
}
Tipos de bloco: "p" paragrafo | "sub" subtopico 2o nivel | "sub3" subtopico 3o nivel |
"b" bullet | "fator" fator SFIT.

Os subtopicos "sub" e "sub3" sao titulos de verdade (Heading 2 e 3), e nao paragrafos em
negrito: e o que permite ao Word montar o sumario com numero de pagina.
"""

try:  # ticket automatico de erro (ver _scripts/erro_ticket.py e a skill /aft-erro)
    import sys as _sys
    from pathlib import Path as _Path
    _aqui = _Path(__file__).resolve()
    for _p in (_aqui.parent, *(_a / "_scripts" for _a in _aqui.parents)):
        if (_p / "erro_ticket.py").is_file():
            _sys.path.insert(0, str(_p))
            from erro_ticket import ativar as _ativar_ticket
            _ativar_ticket(__file__)
            break
except Exception:
    pass

import sys, os, json, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:  # cabecalho institucional com a lotacao do AFT (ver _scripts/cabecalho.py)
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_scripts"))
    from cabecalho import aplicar_no_arquivo
except Exception:  # sem ele o relatorio sai sem cabecalho, como antes
    def aplicar_no_arquivo(caminho, lotacao=None):
        pass

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
PRETO = RGBColor(0x00, 0x00, 0x00)

def add_sumario(doc):
    """Insere o campo TOC do Word (niveis 1 a 3). O Word calcula as paginas ao abrir."""
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('SUMARIO'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL

    p = doc.add_paragraph()
    r = p.add_run()
    ini = OxmlElement('w:fldChar'); ini.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = 'Sumario: clique com o botao direito sobre esta linha e escolha Atualizar campo.'
    fim = OxmlElement('w:fldChar'); fim.set(qn('w:fldCharType'), 'end')
    for el in (ini, instr, sep, txt, fim):
        r._r.append(el)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    try:  # pede ao Word para atualizar os campos na abertura
        el = OxmlElement('w:updateFields'); el.set(qn('w:val'), 'true')
        doc.settings.element.append(el)
    except Exception:
        pass

def add_runs(p, texto, size=11, base_bold=False):
    """Interpreta **negrito** inline e adiciona os runs ao paragrafo."""
    for parte in re.split(r'(\*\*[^*]+\*\*)', texto):
        if parte.startswith('**') and parte.endswith('**'):
            r = p.add_run(parte[2:-2]); r.bold = True
        else:
            r = p.add_run(parte); r.bold = base_bold
        r.font.size = Pt(size)

def main(json_path):
    with open(json_path, 'r', encoding='utf-8') as f:
        d = json.load(f)

    doc = Document()
    nrm = doc.styles['Normal']; nrm.font.name = 'Calibri'; nrm.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3.0); s.right_margin = Cm(2.5)

    def par(texto, bold=False, italic=False, align='just', size=11, after=6, before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = {'just': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                       'left': WD_ALIGN_PARAGRAPH.LEFT}.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        if italic:
            r = p.add_run(texto); r.italic = True; r.font.size = Pt(size); r.bold = bold
        else:
            add_runs(p, texto, size=size, base_bold=bold)
        return p

    # ---- cabecalho ----
    # O cabecalho institucional (brasao, MTE, SIT, lotacao do AFT e logos) entra
    # depois de gravar, pelo _scripts/cabecalho.py; aqui so criamos a parte de
    # cabecalho do arquivo, para ele ter onde entrar.
    for s in doc.sections:
        s.header.is_linked_to_previous = False

    tt = doc.add_paragraph(); tt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tt.add_run(d.get('titulo', 'RELATORIO DE ANALISE DE ACIDENTE DO TRABALHO'))
    r.bold = True; r.font.size = Pt(15); r.font.color.rgb = AZUL
    if d.get('subtitulo'):
        st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = st.add_run(d['subtitulo']); r.italic = True; r.font.size = Pt(11)
    doc.add_paragraph()

    # ---- tabela de identificacao ----
    ident = d.get('identificacao') or []
    if ident:
        tb = doc.add_table(rows=0, cols=2); tb.style = 'Light Grid Accent 1'
        for k, v in ident:
            cells = tb.add_row().cells
            cells[0].paragraphs[0].add_run(str(k)).bold = True
            cells[1].paragraphs[0].add_run(str(v))
            for cc in cells:
                for pp in cc.paragraphs: pp.paragraph_format.space_after = Pt(2)
        doc.add_paragraph()

    if d.get('sumario'):
        add_sumario(doc)

    # ---- secoes ----
    def subtopico(texto, nivel):
        h = doc.add_heading(level=nivel)
        h.paragraph_format.space_before = Pt(8 if nivel == 2 else 6)
        h.paragraph_format.space_after = Pt(2)
        if nivel == 3:
            h.paragraph_format.left_indent = Cm(0.5)
        rr = h.add_run(texto)
        rr.bold = True; rr.font.color.rgb = PRETO
        rr.font.size = Pt(11.5 if nivel == 2 else 11)
        return h

    for sec in d.get('secoes', []):
        h = doc.add_heading(level=1)
        rr = h.add_run(sec.get('titulo', '')); rr.font.color.rgb = AZUL; rr.font.size = Pt(13)
        for bloco in sec.get('blocos', []):
            t = bloco.get('t', 'p')
            if t == 'p':
                par(bloco.get('x', ''))
            elif t == 'sub':
                subtopico(bloco.get('x', ''), 2)
            elif t == 'sub3':
                subtopico(bloco.get('x', ''), 3)
            elif t == 'b':
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                add_runs(p, bloco.get('x', ''))
            elif t == 'fator':
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(4); p.paragraph_format.line_spacing = 1.15
                rr = p.add_run('Fator Causal %s - %s' % (bloco.get('codigo', ''), bloco.get('nome', '')))
                rr.bold = True; rr.font.size = Pt(11)
                p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p2.paragraph_format.space_after = Pt(6); p2.paragraph_format.line_spacing = 1.15
                lab = p2.add_run('Classificacao (%s). Descricao: ' % bloco.get('classe', 'a confirmar'))
                lab.bold = True; lab.font.size = Pt(11)
                add_runs(p2, bloco.get('desc', ''))

    if d.get('rodape'):
        doc.add_paragraph()
        par(d['rodape'], italic=True, size=9, after=0)

    saida = d['saida']
    os.makedirs(os.path.dirname(saida), exist_ok=True) if os.path.dirname(saida) else None
    doc.save(saida)
    aplicar_no_arquivo(saida)
    print('SALVO:', saida)

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Uso: python gerar_relatorio_docx.py conteudo.json'); sys.exit(1)
    main(sys.argv[1])
